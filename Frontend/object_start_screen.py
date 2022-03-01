# Class for Start Screen Frame

from pathlib import Path
from tkinter.ttk import Frame, Entry, Button, Combobox
from tkinter import Tk, Canvas, PhotoImage
from tkinter.messagebox import showwarning
import pptx
import openpyxl
from Frontend.commands import *
from Frontend.strings import *
from Backend.Util_Functions import *

# When class is initialized, the startscreen will create a frame and canvas with relevant widgets
class StartScreenFrame():
    def __init__(self,root: Tk,controller):
        super().__init__()

        # Locates the assests folder and stores it in ASSETS_PATH
        self.OUTPUT_PATH = Path(__file__).parent
        self.ASSETS_PATH = self.OUTPUT_PATH / Path(str_assets)

        # Initialises the root and controller classes
        self.controller = controller
        self.root=root

        # Creates a frame in the root window and places it such that it covers whole window
        self.frame = Frame(root)
        self.frame.grid(row=0,column=0,sticky="nsew")
        self.frame.configure(bg = "#938E8E")

        # Creates empty list used to fill column names values
        self.ganttData_column_names=[]

        # Creates a canvas with desired properties
        self.canvas = Canvas(
            self.frame,
            bg = "#938E8E",
            height = 700,
            width = 984,
            bd = 0,
            highlightthickness = 0,
            relief = "ridge"
        )
        # Places a rectange with desired color to cover whole canvas
        self.canvas.place(x = 0, y = 0)
        self.canvas.create_rectangle(
            0.0,
            0.0,
            984.0,
            700.0,
            fill="#2E2E2E",
            outline="")

        # Title
        self.canvas.create_text(
            283.0,
            42.0,
            anchor="nw",
            text=str_project_title,
            fill="#FFFFFF",
            font=("RobotoRoman SemiBold", 48 * -1)
        )

        # Required Labels
        self.canvas.create_text(
            54.0,
            204.0,
            anchor="nw",
            text=str_select_powerpoint_file,
            fill="#FFFFFF",
            font=("RobotoRoman Medium", 30 * -1)
        )

        self.canvas.create_text(
            54.0,
            368.0,
            anchor="nw",
            text=str_select_data,
            fill="#FFFFFF",
            font=("RobotoRoman Medium", 30 * -1)
        )
        
        self.canvas.create_text(
            54.0,
            500.0,
            anchor="nw",
            text=str_select_sheet,
            fill="#FFFFFF",
            font=("RobotoRoman Medium", 30 * -1)
        )

        # Entry for ppt path
        self.entry_presentation_path = Entry(self.frame,
            bd=0,
            bg="#ECECEC",
            highlightthickness=2,
            width=39,
            font=("RobotoRoman",24)
        )
        self.canvas.create_window(54,250,anchor="nw",window=self.entry_presentation_path)
        self.entry_presentation_path.insert(0,str_default_browse_text)

        # Entry for data path
        self.entry_data_path = Entry(self.frame,
            bd=0,
            bg="#ECECEC",
            highlightthickness=2,
            width=39,
            font=("RobotoRoman",24)
        )
        self.canvas.create_window(54,420,anchor="nw",window=self.entry_data_path)
        self.entry_data_path.insert(0,str_default_browse_text)

        # Button to browse ppt path
        self.button_browse_img = PhotoImage(
            file=self.relative_to_assets(str_browse_button_img_name))
        self.button_ppt_browse = Button(self.canvas,
            image=self.button_browse_img,
            borderwidth=0,
            highlightthickness=2,
            command=lambda: browsePowerPointButton(self.root,self.entry_presentation_path),
            relief="raised"
        )
        self.button_ppt_browse.place(
            x=778.0,
            y=243.0,
            width=150.0,
            height=50.0
        )

        # Button to browse data path
        self.button_data_browse = Button(self.canvas,
            image=self.button_browse_img,
            borderwidth=0,
            highlightthickness=2,
            command=lambda: self.self_browse_data_button(),
            relief="raised"
        )
        self.button_data_browse.place(
            x=778.0,
            y=414.0,
            width=150.0,
            height=50.0
        )

        # Button to go to next frame
        self.button_next_img = PhotoImage(
            file=self.relative_to_assets(str_next_button_img_name))
        self.button_next = Button(self.canvas,
            image=self.button_next_img,
            borderwidth=0,
            highlightthickness=2,
            command=lambda: self.self_show_frame(),
            relief="raised"
        )
        self.button_next.place(
            x=417.0,
            y=585.0,
            width=150.0,
            height=50.0
        )
# Function to return relative path
    def relative_to_assets(self,path: str) -> Path:
        return self.ASSETS_PATH / Path(path)

# Function to get path from the input textbox for presentation
    def get_presentation_path(self):
        return self.entry_presentation_path.get()

# Function to get path from the input textbox for data
    def get_data_path(self):
        return self.entry_data_path.get()

# Function to show next frame and send data for further input or importing as dataFrame
    def self_show_frame(self):
        data_path = str(self.get_data_path())
        prs_path = str(self.get_presentation_path())
        try:
            self.ganttPrs = pptx.Presentation(prs_path)
            if ".csv" in data_path:
                self.ganttData = importDataCsv(data_path)
            else:
                sheet_name = str(self.selected_sheet_name.get())
                self.ganttData= importDataExcel(data_path,sheet_name)
            self.ganttData_column_names = self.ganttData.columns.values.tolist()
            showFrame(self.controller.myDataSetup.frame)
            self.controller.myDataSetup.drop_down_columns(self.ganttData_column_names,self.ganttData)
        except:
            showwarning("Data or Presentation Not Selected","Please selet path for data and presentation by clicking on browse")

# Function to show dropdown based on the type of data selected in the browse data
    def self_browse_data_button(self):
        browseDataButton(self.root,self.entry_data_path)
        data_path = str(self.get_data_path())
        if ".csv" in data_path:
            pass
        else:
            wkbook = openpyxl.load_workbook(data_path)
            self.sheet_names=wkbook.sheetnames
            # Dropdown menu with the names of the sheets
            self.selected_sheet_name = StringVar()
            self.comboBox_select_sheet_name = Combobox(
                self.canvas,
                value=self.sheet_names,
                textvariable=self.selected_sheet_name,
                width=109
            )
            self.comboBox_select_sheet_name.current(0)
            self.comboBox_select_sheet_name.state(["readonly"])
            self.canvas.create_window(250,508,anchor="nw",window=self.comboBox_select_sheet_name)

